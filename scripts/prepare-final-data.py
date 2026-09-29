#!/usr/bin/env python3
"""Prepare the TCVA Awards 2026 final workbook for Supabase import.

The generated SQL contains personal guest data and is intentionally written to
an ignored local outputs folder. Do not commit generated outputs.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import uuid
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_WORKBOOK = Path("/workspace/scratch/3a83cbc6d080/upload/Invitation List.xlsx")
DEFAULT_OUTDIR = Path("outputs/final-data")
PROGRAMME_URL = "https://www.tvawales.org.uk/web/content/3696?unique=a3c76b86fb5669f9bb86ffb4eb71d8ec696c1c28&download=true"
NAMESPACE = uuid.UUID("6f166831-bbcb-4fb4-a995-5fe87f48c0a8")

DIETARY_WORDS = re.compile(
    r"\b(vegetarian|veggie|vegan|gluten|celiac|coeliac|allerg|dairy|fish|seafood|nuts?|peanuts?|"
    r"onion|mushroom|mustard|sesame|chicken|pork|turkey|cheese|lactose|no food|no main|free)\b",
    re.I,
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def norm_name(value: str) -> str:
    value = clean(value).lower()
    value = value.replace("ll ", "")
    value = re.sub(r"\([^)]*\)", "", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    value = re.sub(r"\b(no|none|na|n a)\b", "", value)
    return re.sub(r"\s+", " ", value).strip()


def sql_text(value: Any) -> str:
    if value is None:
        return "null"
    value = clean(value)
    if value == "":
        return "null"
    return "'" + value.replace("'", "''") + "'"


def sql_bool(value: bool) -> str:
    return "true" if value else "false"


def stable_uuid(kind: str, key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f"{kind}:{key}"))


def parse_table_number(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    match = re.search(r"(\d+)", text)
    if not match:
        return ""
    number = int(match.group(1))
    if number <= 0:
        return ""
    return str(number)


def confirmation_from(value: Any) -> str:
    text = clean(value).lower()
    if text in {"y", "yes"}:
        return "confirmed"
    if text in {"n", "no"}:
        return "declined"
    return "tbc"


def looks_like_name(value: str) -> bool:
    text = clean(value)
    if not text:
        return False
    lowered = text.lower()
    if lowered in {"no", "none", "n/a", "na", "no +1", "no plus one"}:
        return False
    if DIETARY_WORDS.search(text) and len(re.findall(r"[A-Z][a-z]+", text)) < 1:
        return False
    return bool(re.search(r"[A-Za-z]", text))


def split_people(value: Any) -> list[tuple[str, str]]:
    text = clean(value)
    if not looks_like_name(text):
        return []

    text = text.replace(" and ", " & ")
    lines = [line.strip() for line in re.split(r"\n|;", str(value)) if clean(line)]
    if not lines:
        lines = [text]

    people: list[tuple[str, str]] = []
    for line in lines:
        line = clean(line)
        if not looks_like_name(line):
            continue
        if " - " in line:
            name, note = line.split(" - ", 1)
            if looks_like_name(name):
                people.append((clean(name), clean(note)))
            continue
        parts = [clean(part) for part in re.split(r"\s*&\s*|,\s*", line) if clean(part)]
        if len(parts) > 1 and all(looks_like_name(part) for part in parts):
            people.extend((part, "") for part in parts)
        else:
            people.append((line, ""))
    return people


@dataclass
class Candidate:
    key: str
    full_name: str
    party_key: str
    source_sheet: str
    source_row: int
    invite_number: str = ""
    organisation_name: str = ""
    role_label: str = ""
    relationship_label: str = ""
    award_category: str = ""
    sponsor_name: str = ""
    phone: str = ""
    email: str = ""
    dietary_notes: str = ""
    accessibility_notes: str = ""
    linked_notes: str = ""
    table_number: str = ""
    confirmation_status: str = "tbc"
    is_placeholder: bool = False
    visual_table_number: str = ""
    seat_sort_order: int | None = None
    admin_notes: list[str] = field(default_factory=list)

    def merge_from(self, other: "Candidate") -> None:
        for attr in [
            "organisation_name",
            "role_label",
            "relationship_label",
            "award_category",
            "sponsor_name",
            "phone",
            "email",
            "dietary_notes",
            "accessibility_notes",
            "linked_notes",
            "invite_number",
        ]:
            if not getattr(self, attr) and getattr(other, attr):
                setattr(self, attr, getattr(other, attr))
        if self.confirmation_status == "tbc" and other.confirmation_status != "tbc":
            self.confirmation_status = other.confirmation_status
        if other.admin_notes:
            self.admin_notes.extend(other.admin_notes)


def load_table_plan(ws) -> list[Candidate]:
    headers: list[tuple[int, int, str]] = []
    for row in ws.iter_rows():
        for cell in row:
            table_number = parse_table_number(cell.value)
            if table_number and "table" in clean(cell.value).lower():
                headers.append((cell.row, cell.column, table_number))

    visual: list[Candidate] = []
    for header_row, column, table_number in headers:
        next_header = min(
            [row for row, col, _num in headers if col == column and row > header_row] or [ws.max_row + 1]
        )
        order = 0
        for row_idx in range(header_row + 1, next_header):
            raw_name = clean(ws.cell(row_idx, column).value)
            if not raw_name:
                continue
            order += 1
            name = raw_name
            accessibility = ""
            if re.search(r"wheelchair", name, re.I):
                accessibility = "Wheelchair"
                name = clean(re.sub(r"\([^)]*wheelchair[^)]*\)", "", name, flags=re.I))
            is_placeholder = bool(re.search(r"\bx\s*\d+\b", name, re.I))
            key = f"visual:{table_number}:{order}:{name}"
            visual.append(
                Candidate(
                    key=key,
                    full_name=name,
                    party_key=key,
                    source_sheet="Table Plan",
                    source_row=row_idx,
                    visual_table_number=table_number,
                    table_number=table_number,
                    seat_sort_order=order,
                    accessibility_notes=accessibility,
                    confirmation_status="confirmed",
                    is_placeholder=is_placeholder,
                )
            )
    return visual


def source_candidates(wb) -> list[Candidate]:
    candidates: list[Candidate] = []
    current_category = ""

    def add_candidate(**kwargs: Any) -> None:
        name = clean(kwargs.get("full_name"))
        if not looks_like_name(name):
            return
        key = f"{kwargs['source_sheet']}:{kwargs['source_row']}:{kwargs.get('role_label','')}:{name}"
        candidates.append(Candidate(key=key, full_name=name, **{k: v for k, v in kwargs.items() if k != "full_name"}))

    ws = wb["Winners, Runner Ups & Nominee"]
    for row in range(2, ws.max_row + 1):
        invite = clean(ws.cell(row, 1).value)
        category = clean(ws.cell(row, 2).value)
        if category:
            current_category = category
        relationship = clean(ws.cell(row, 3).value)
        primary = clean(ws.cell(row, 4).value)
        phone = clean(ws.cell(row, 6).value)
        address = clean(ws.cell(row, 7).value)
        comments = clean(ws.cell(row, 8).value)
        attending = confirmation_from(ws.cell(row, 11).value)
        guest_cell = ws.cell(row, 12).value
        dietary = clean(ws.cell(row, 13).value)
        accessibility = clean(ws.cell(row, 14).value)
        table = parse_table_number(ws.cell(row, 15).value)
        party_key = f"winners:{row}:{invite or primary}"
        note = "; ".join(part for part in [comments, address] if part)
        add_candidate(
            party_key=party_key,
            source_sheet=ws.title,
            source_row=row,
            invite_number=invite,
            full_name=primary,
            role_label=relationship,
            relationship_label=relationship,
            award_category=current_category,
            phone=phone,
            dietary_notes="" if dietary.lower() in {"no", "none", "na", "n/a"} else dietary,
            accessibility_notes="" if accessibility.lower() in {"no", "none", "na", "n/a"} else accessibility,
            linked_notes=note,
            table_number=table,
            confirmation_status=attending,
        )
        for index, (guest_name, guest_note) in enumerate(split_people(guest_cell), start=1):
            add_candidate(
                party_key=party_key,
                source_sheet=ws.title,
                source_row=row,
                invite_number=invite,
                full_name=guest_name,
                role_label="Linked guest",
                relationship_label="Guest",
                award_category=current_category,
                phone=phone,
                dietary_notes=guest_note,
                accessibility_notes="",
                linked_notes=f"Linked to {primary}".strip(),
                table_number=table,
                confirmation_status=attending,
            )

    ws = wb["Judges, Sponsors"]
    for row in range(2, ws.max_row + 1):
        role = clean(ws.cell(row, 2).value)
        organisation = clean(ws.cell(row, 3).value)
        primary = clean(ws.cell(row, 4).value)
        attending = confirmation_from(ws.cell(row, 9).value)
        dietary = clean(ws.cell(row, 11).value)
        table = parse_table_number(ws.cell(row, 13).value)
        party_key = f"judges-sponsors:{row}:{role or primary}"
        add_candidate(
            party_key=party_key,
            source_sheet=ws.title,
            source_row=row,
            full_name=primary,
            organisation_name=organisation,
            role_label=role,
            relationship_label=role,
            dietary_notes="" if dietary.lower() in {"no", "none", "na", "n/a"} else dietary,
            table_number=table,
            confirmation_status=attending,
        )
        for guest_name, guest_note in split_people(ws.cell(row, 10).value):
            add_candidate(
                party_key=party_key,
                source_sheet=ws.title,
                source_row=row,
                full_name=guest_name,
                organisation_name=organisation,
                role_label="Linked guest",
                relationship_label="Guest",
                dietary_notes=guest_note,
                linked_notes=f"Linked to {primary}".strip(),
                table_number=table,
                confirmation_status=attending,
            )

    ws = wb["Staff & Comms"]
    for row in range(2, ws.max_row + 1):
        role = clean(ws.cell(row, 1).value)
        organisation = clean(ws.cell(row, 2).value)
        primary = clean(ws.cell(row, 3).value) or role
        attending = confirmation_from(ws.cell(row, 6).value)
        dietary = clean(ws.cell(row, 8).value)
        table = parse_table_number(ws.cell(row, 10).value)
        party_key = f"staff:{row}:{role or primary}"
        add_candidate(
            party_key=party_key,
            source_sheet=ws.title,
            source_row=row,
            full_name=primary,
            organisation_name=organisation,
            role_label=role,
            relationship_label=role,
            dietary_notes="" if dietary.lower() in {"no", "none", "na", "n/a"} else dietary,
            table_number=table,
            confirmation_status=attending,
        )
        for guest_name, guest_note in split_people(ws.cell(row, 7).value):
            add_candidate(
                party_key=party_key,
                source_sheet=ws.title,
                source_row=row,
                full_name=guest_name,
                organisation_name=organisation,
                role_label="Linked guest",
                relationship_label="Guest",
                dietary_notes=guest_note,
                linked_notes=f"Linked to {primary}".strip(),
                table_number=table,
                confirmation_status=attending,
            )

    ws = wb["Waiting List"]
    for row in range(2, ws.max_row + 1):
        name = clean(ws.cell(row, 1).value)
        if not name:
            continue
        add_candidate(
            party_key=f"waitlist:{row}:{name}",
            source_sheet=ws.title,
            source_row=row,
            full_name=name,
            organisation_name=clean(ws.cell(row, 2).value),
            role_label="Waiting list",
            relationship_label="Waiting list",
            linked_notes=f"Tickets requested: {clean(ws.cell(row, 3).value)}. {clean(ws.cell(row, 4).value)}",
            confirmation_status="waitlist",
        )

    return candidates


def dietary_lookup(wb) -> dict[str, str]:
    ws = wb["Dietary Requirements"]
    lookup: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        name = clean(ws.cell(row, 2).value)
        dietary = clean(ws.cell(row, 3).value)
        notes = clean(ws.cell(row, 4).value)
        combined = "; ".join(part for part in [dietary, notes] if part)
        if name and combined:
            lookup[norm_name(name)] = combined
    return lookup


def best_match(candidate: Candidate, visual_index: dict[str, Candidate], all_visual: list[Candidate]) -> Candidate | None:
    exact = visual_index.get(norm_name(candidate.full_name))
    if exact:
        return exact
    cleaned = re.split(r"\s+-\s+|\s+\(", candidate.full_name, 1)[0]
    exact = visual_index.get(norm_name(cleaned))
    if exact:
        return exact
    wanted = norm_name(cleaned)
    if not wanted:
        return None
    scored = [
        (SequenceMatcher(None, wanted, norm_name(v.full_name)).ratio(), v)
        for v in all_visual
        if abs(len(wanted) - len(norm_name(v.full_name))) <= 8
    ]
    if not scored:
        return None
    score, match = max(scored, key=lambda item: item[0])
    return match if score >= 0.88 else None


def build_records(workbook: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    wb = load_workbook(workbook, data_only=True)
    visual = load_table_plan(wb["Table Plan"])
    structured = source_candidates(wb)
    diet = dietary_lookup(wb)
    warnings: list[str] = []

    visual_index: dict[str, Candidate] = {}
    for visual_guest in visual:
        visual_index.setdefault(norm_name(visual_guest.full_name), visual_guest)

    matched_visual_keys: set[str] = set()
    final: list[Candidate] = []

    for visual_guest in visual:
        lookup = diet.get(norm_name(visual_guest.full_name))
        if lookup and not visual_guest.dietary_notes:
            visual_guest.dietary_notes = lookup
        final.append(visual_guest)

    final_by_visual_key = {guest.key: guest for guest in final}
    for structured_guest in structured:
        lookup = diet.get(norm_name(structured_guest.full_name))
        if lookup:
            structured_guest.dietary_notes = lookup
        visual_match = best_match(structured_guest, visual_index, visual)
        if visual_match:
            target = final_by_visual_key[visual_match.key]
            target.merge_from(structured_guest)
            matched_visual_keys.add(visual_match.key)
            if structured_guest.table_number and structured_guest.table_number != target.visual_table_number:
                target.admin_notes.append(
                    f"Structured source table {structured_guest.table_number}; visual table plan uses {target.visual_table_number}."
                )
            continue

        include = structured_guest.confirmation_status in {"confirmed", "waitlist"} or structured_guest.table_number
        if include:
            if structured_guest.table_number and int(structured_guest.table_number) > 20:
                warning = (
                    f"{structured_guest.source_sheet} row {structured_guest.source_row}: "
                    f"{structured_guest.full_name} has unusual table number {structured_guest.table_number}."
                )
                warnings.append(warning)
                structured_guest.admin_notes.append(warning)
                structured_guest.table_number = ""
            final.append(structured_guest)

    tables: dict[str, dict[str, Any]] = {}
    for guest in final:
        table_number = guest.visual_table_number or guest.table_number
        if table_number:
            tables.setdefault(
                table_number,
                {
                    "id": stable_uuid("table", table_number),
                    "table_number": table_number,
                    "table_label": f"Table {table_number}",
                    "capacity": 10,
                    "display_order": int(table_number),
                },
            )

    parties: dict[str, dict[str, Any]] = {}
    guests: list[dict[str, Any]] = []
    seen_names: dict[str, int] = {}
    for guest in final:
        name_key = norm_name(guest.full_name) or guest.full_name.lower()
        seen_names[name_key] = seen_names.get(name_key, 0) + 1
        dedupe_suffix = seen_names[name_key]
        guest_id = stable_uuid("guest", f"{guest.source_sheet}:{guest.source_row}:{guest.full_name}:{dedupe_suffix}")
        party_id = stable_uuid("party", guest.party_key)
        parties.setdefault(
            guest.party_key,
            {
                "id": party_id,
                "party_code": guest.party_key[:80],
                "party_name": guest.full_name,
                "lead_guest_name": guest.full_name,
                "organisation_name": guest.organisation_name,
                "award_category": guest.award_category,
                "sponsor_name": guest.sponsor_name,
                "notes": "",
            },
        )
        table_number = guest.visual_table_number or guest.table_number
        guests.append(
            {
                "id": guest_id,
                "party_id": party_id,
                "table_id": tables[table_number]["id"] if table_number else None,
                "invite_number": guest.invite_number,
                "source_sheet": guest.source_sheet,
                "source_row_number": guest.source_row,
                "full_name": guest.full_name,
                "organisation_name": guest.organisation_name,
                "role_label": guest.role_label,
                "relationship_label": guest.relationship_label,
                "award_category": guest.award_category,
                "sponsor_name": guest.sponsor_name,
                "phone": guest.phone,
                "email": guest.email,
                "dietary_notes": guest.dietary_notes,
                "accessibility_notes": guest.accessibility_notes,
                "confirmation_status": guest.confirmation_status,
                "is_placeholder": guest.is_placeholder,
                "linked_notes": guest.linked_notes,
                "seat_sort_order": guest.seat_sort_order,
                "seating_status": "assigned" if table_number else ("not_required" if guest.confirmation_status == "declined" else "unassigned"),
                "admin_notes": "; ".join(dict.fromkeys(note for note in guest.admin_notes if note)),
            }
        )

    return list(tables.values()), list(parties.values()), guests, warnings


def generate_sql(workbook: Path, tables: list[dict[str, Any]], parties: list[dict[str, Any]], guests: list[dict[str, Any]], warnings: list[str]) -> str:
    digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
    batch_id = stable_uuid("batch", digest)
    rows: list[str] = [
        "-- TCVA Awards 2026 final event data import",
        "-- Generated locally from Invitation List.xlsx. Contains guest data; do not commit.",
        "begin;",
        "",
        "insert into public.import_batches (id, source_file_name, source_file_hash, rows_seen, rows_imported, warnings)",
        f"values ('{batch_id}', {sql_text(workbook.name)}, '{digest}', {len(guests)}, {len(guests)}, '{json.dumps(warnings).replace(chr(39), chr(39)+chr(39))}'::jsonb)",
        "on conflict (id) do nothing;",
        "",
        "insert into public.event_assets (asset_key, label, asset_type, url, qr_enabled, active, display_order, notes)",
        f"values ('programme_pdf', 'Digital programme PDF', 'programme_pdf', {sql_text(PROGRAMME_URL)}, true, true, 1, 'Final print programme download link')",
        "on conflict (asset_key) do update set url = excluded.url, active = true, qr_enabled = true, updated_at = now();",
        "",
    ]

    if tables:
        rows.append("insert into public.event_tables (id, table_number, table_label, capacity, display_order, active)")
        values = [
            f"('{t['id']}', {sql_text(t['table_number'])}, {sql_text(t['table_label'])}, {t['capacity']}, {t['display_order']}, true)"
            for t in sorted(tables, key=lambda item: item["display_order"])
        ]
        rows.append("values\n  " + ",\n  ".join(values))
        rows.append(
            "on conflict (table_number) do update set table_label = excluded.table_label, capacity = excluded.capacity, display_order = excluded.display_order, active = true, updated_at = now();"
        )
        rows.append("")

    if parties:
        rows.append("insert into public.guest_parties (id, party_code, party_name, lead_guest_name, organisation_name, award_category, sponsor_name, notes)")
        values = [
            "("
            + ", ".join(
                [
                    f"'{p['id']}'",
                    sql_text(p["party_code"]),
                    sql_text(p["party_name"]),
                    sql_text(p["lead_guest_name"]),
                    sql_text(p["organisation_name"]),
                    sql_text(p["award_category"]),
                    sql_text(p["sponsor_name"]),
                    sql_text(p["notes"]),
                ]
            )
            + ")"
            for p in parties
        ]
        rows.append("values\n  " + ",\n  ".join(values))
        rows.append("on conflict (party_code) do nothing;")
        rows.append("")

    if guests:
        columns = [
            "id",
            "party_id",
            "table_id",
            "import_batch_id",
            "invite_number",
            "source_sheet",
            "source_row_number",
            "full_name",
            "organisation_name",
            "role_label",
            "relationship_label",
            "award_category",
            "sponsor_name",
            "email",
            "phone",
            "dietary_notes",
            "accessibility_notes",
            "confirmation_status",
            "is_placeholder",
            "linked_notes",
            "seat_sort_order",
            "seating_status",
            "admin_notes",
        ]
        rows.append(f"insert into public.guests ({', '.join(columns)})")
        values = []
        for guest in guests:
            values.append(
                "("
                + ", ".join(
                    [
                        f"'{guest['id']}'",
                        f"'{guest['party_id']}'",
                        f"'{guest['table_id']}'" if guest["table_id"] else "null",
                        f"'{batch_id}'",
                        sql_text(guest["invite_number"]),
                        sql_text(guest["source_sheet"]),
                        str(guest["source_row_number"]),
                        sql_text(guest["full_name"]),
                        sql_text(guest["organisation_name"]),
                        sql_text(guest["role_label"]),
                        sql_text(guest["relationship_label"]),
                        sql_text(guest["award_category"]),
                        sql_text(guest["sponsor_name"]),
                        sql_text(guest["email"]),
                        sql_text(guest["phone"]),
                        sql_text(guest["dietary_notes"]),
                        sql_text(guest["accessibility_notes"]),
                        f"'{guest['confirmation_status']}'::public.guest_confirmation_status",
                        sql_bool(guest["is_placeholder"]),
                        sql_text(guest["linked_notes"]),
                        str(guest["seat_sort_order"]) if guest["seat_sort_order"] is not None else "null",
                        sql_text(guest["seating_status"]),
                        sql_text(guest["admin_notes"]),
                    ]
                )
                + ")"
            )
        rows.append("values\n  " + ",\n  ".join(values))
        rows.append("on conflict (id) do nothing;")
        rows.append("")

    rows.append("commit;")
    rows.append("")
    rows.append("select 'event_tables' as table_name, count(*) from public.event_tables")
    rows.append("union all select 'guest_parties', count(*) from public.guest_parties")
    rows.append("union all select 'guests', count(*) from public.guests")
    rows.append("union all select 'event_assets', count(*) from public.event_assets;")
    return "\n".join(rows)


def write_report(outdir: Path, tables: list[dict[str, Any]], parties: list[dict[str, Any]], guests: list[dict[str, Any]], warnings: list[str]) -> None:
    assigned = sum(1 for guest in guests if guest["table_id"])
    dietary = sum(1 for guest in guests if guest["dietary_notes"])
    placeholders = sum(1 for guest in guests if guest["is_placeholder"])
    waitlist = sum(1 for guest in guests if guest["confirmation_status"] == "waitlist")
    report = [
        "# TCVA Awards 2026 Import Validation Report",
        "",
        f"- Tables prepared: {len(tables)}",
        f"- Guest parties prepared: {len(parties)}",
        f"- Guest records prepared: {len(guests)}",
        f"- Guests assigned to tables: {assigned}",
        f"- Guests with dietary notes: {dietary}",
        f"- Placeholder/group records: {placeholders}",
        f"- Waiting list records: {waitlist}",
        "",
        "## Warnings",
        "",
    ]
    if warnings:
        report.extend(f"- {warning}" for warning in warnings)
    else:
        report.append("- No high-priority import warnings.")
    (outdir / "validation_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--outdir", type=Path, default=DEFAULT_OUTDIR)
    parser.add_argument("--summary-only", action="store_true")
    args = parser.parse_args()

    tables, parties, guests, warnings = build_records(args.workbook)
    summary = {
        "tables": len(tables),
        "parties": len(parties),
        "guests": len(guests),
        "assigned": sum(1 for guest in guests if guest["table_id"]),
        "dietary": sum(1 for guest in guests if guest["dietary_notes"]),
        "placeholders": sum(1 for guest in guests if guest["is_placeholder"]),
        "warnings": warnings,
    }
    print(json.dumps(summary, indent=2))

    if args.summary_only:
        return

    args.outdir.mkdir(parents=True, exist_ok=True)
    (args.outdir / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (args.outdir / "records.json").write_text(json.dumps({"tables": tables, "parties": parties, "guests": guests}, indent=2), encoding="utf-8")
    (args.outdir / "tcva_awards_2026_import.sql").write_text(
        generate_sql(args.workbook, tables, parties, guests, warnings),
        encoding="utf-8",
    )
    write_report(args.outdir, tables, parties, guests, warnings)


if __name__ == "__main__":
    main()
